#!/usr/bin/env python3
"""Build the implementation chapter and exact, data-backed figures.
Run from repo root: .venv/bin/python "product analytics/build_report.py"
Dependencies: requirements.txt in this directory. No network or live services used.
"""
import csv
import json
import math
import sys
import textwrap
import zipfile
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path
from html import escape

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
import markdown

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
DATA = ROOT / 'analytics/data'
IMG = OUT / 'images'
IMG.mkdir(exist_ok=True)
ASOF = date(2026, 9, 20)
def read(name):
    with (DATA / (name + '.csv')).open(newline='') as f:
        return list(csv.DictReader(f))
users = read('user_signups')
sessions = read('monitoring_sessions')
configs = read('api_key_configs')
visitors = read('web_visitors')
monthly = read('kpi_monthly')
summary = json.loads((ROOT / 'analytics/summary.json').read_text())
activated = [u for u in users if u['activated_within_7d_flag'] == '1']
inactive = [u for u in activated if u['inactive_30d_flag'] == '1']
session_dates = defaultdict(set)
for s in sessions:
    session_dates[s['user_id']].add(date.fromisoformat(s['session_date']))
observed = len(session_dates)
missing = sum(not session_dates[u['user_id']] for u in activated)
# Full follow-up only; do not count incomplete week-four windows as failures.
w4 = []
for month in sorted({u['cohort_month'] for u in activated}):
    eligible = [u for u in activated if u['cohort_month'] == month and date.fromisoformat(u['activation_date']) + timedelta(days=35) <= ASOF]
    retained = sum(any(28 <= (d - date.fromisoformat(u['activation_date'])).days < 35 for d in session_dates[u['user_id']]) for u in eligible)
    if eligible:
        w4.append({'month': month, 'eligible': len(eligible), 'retained': retained, 'rate': round(100 * retained / len(eligible), 1)})

# Audit invariants without editing the repository's original datasets.
assert len({u['user_id'] for u in users}) == len(users)
assert len({v['visitor_id'] for v in visitors}) == len(visitors)
assert len({s['session_id'] for s in sessions}) == len(sessions)
user_ids = {u['user_id'] for u in users}
config_ids = {c['config_id'] for c in configs}
assert all(s['user_id'] in user_ids and s['config_id'] in config_ids for s in sessions)
assert sum(int(u['total_sessions']) for u in users) == len(sessions)
assert len(users) == summary['n_signups'] and len(activated) == summary['n_activated']

NAVY = '#142b49'; TEAL = '#008c95'; BLUE = '#387ed1'; ORANGE = '#df8735'; RED = '#bd4e55'; GRAY = '#607387'
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 11, 'axes.spines.top': False, 'axes.spines.right': False, 'axes.labelcolor': NAVY, 'text.color': NAVY, 'xtick.color': GRAY, 'ytick.color': NAVY, 'figure.facecolor': '#ffffff', 'axes.facecolor': '#ffffff', 'savefig.facecolor': '#ffffff'})

def finish(fig, name, title, note='Synthetic analytics | 01 Mar - 20 Sep 2026 | Source: repository CSVs'):
    fig.suptitle(title, x=.055, ha='left', fontsize=18, fontweight='bold', color=NAVY)
    fig.text(.055, .025, note, fontsize=9, color=GRAY)
    fig.tight_layout(rect=[.025,.07,.98,.90])
    fig.savefig(IMG / name, dpi=180)
    plt.close(fig)

def diagram(name, title, nodes, note, colors_=None):
    fig, ax = plt.subplots(figsize=(12,5))
    ax.set_xlim(0,12); ax.set_ylim(0,5); ax.axis('off')
    n=len(nodes); width=2.45 if n<=4 else 1.95; step=11.6/n
    for i,(heading,body) in enumerate(nodes):
        x=.15+i*step; y=1.25
        c=(colors_ or [TEAL]*n)[i]
        ax.add_patch(FancyBboxPatch((x,y),width,2.4,boxstyle='round,pad=0.08,rounding_size=0.12',facecolor='#edf5fa',edgecolor=c,linewidth=1.6))
        ax.text(x+width/2,3.12,heading,ha='center',va='center',fontweight='bold',fontsize=12,color=c)
        ax.text(x+width/2,2.2,body,ha='center',va='center',fontsize=10,linespacing=1.65)
        if i<n-1:
            ax.annotate('',xy=(x+step-.12,2.5),xytext=(x+width+.11,2.5),arrowprops={'arrowstyle':'->','color':GRAY,'lw':2})
    finish(fig,name,title,note)

# Figure 4.1
fig,ax=plt.subplots(figsize=(12,5))
counts=[len(visitors),len(users),summary['n_config_users'],len(activated)]
labels=['Website visitors','Account signups','Connected-key users','Activation flags']
ax.barh(labels[::-1],counts[::-1],color=[TEAL,BLUE,ORANGE,NAVY][::-1],height=.6)
for i,v in enumerate(counts[::-1]):
    ax.text(v+240,i,f'{v:,}  |  {v/counts[0]*100:.2f}% of visitors',va='center',fontsize=11)
ax.set_xlim(0,28000); ax.set_xlabel('Unique users / visitors, not configuration rows')
ax.grid(axis='x',alpha=.15); ax.set_axisbelow(True)
finish(fig,'4_1_funnel.png','4.1  |  Acquisition and activation funnel')
# Figure 4.2
fig,ax=plt.subplots(figsize=(12,5))
locs=['All activated','DSU Campus','Outside Campus']
nums=[(295,813),(81,226),(214,587)]
rates=[n/d*100 for n,d in nums]
ax.barh(locs[::-1],rates[::-1],color=[NAVY,TEAL,BLUE][::-1],height=.55)
for i,(r,(n,d)) in enumerate(zip(rates[::-1],nums[::-1])):
    ax.text(r+1,i,f'{r:.1f}%  ({n}/{d})',va='center')
ax.set_xlim(0,55); ax.set_xlabel('More than 30 days since last session (%)')
finish(fig,'4_2_inactivity.png','4.2  |  Inactivity proxy, not subscription churn','Synthetic data | Snapshot: 20 Sep 2026 | Denominator: all activation flags; not exposure-adjusted')
# Figure 4.3
diagram('4_3_import_pipeline.png','4.3  |  Data generation and import architecture',[
    ('LIVE SOURCES','GA4 + auth\nConvex telemetry\nIntended sources'),
    ('ETL PATH','Convex export\nWebhook -> Postgres\nSchema validation'),
    ('TRANSFORM','dbt staging\nDaily / provider marts\nQuality metrics'),
    ('ANALYTICS','Live marts (intended)\nSeparate synthetic CSVs\nTableau + report')],
    'Conceptual source-to-analysis map | CSV analytics pack is separate from the live ETL; deployment not verified')
# Figure 4.4
fig,ax=plt.subplots(figsize=(12,5))
loc=['DSU Campus','Outside Campus']; x=np.arange(2)
a=[summary['location'][l]['activation_rate'] for l in loc]
b=[summary['location'][l]['inactive_rate'] for l in loc]
ax.bar(x-.18,a,.34,color=TEAL,label='Activated / signups (%)')
ax.bar(x+.18,b,.34,color=ORANGE,label='Inactive / activated (%)')
for xs,vals in [(x-.18,a),(x+.18,b)]:
    for xx,v in zip(xs,vals): ax.text(xx,v+1,f'{v:.1f}%',ha='center')
ax.set_xticks(x,loc);ax.set_ylim(0,80);ax.set_ylabel('Percentage, denominators differ');ax.legend(loc='upper left',frameon=False)
finish(fig,'4_4_segment_evidence.png','4.4  |  An activation gap, not a campus churn gap')
# Figure 4.5
diagram('4_5_strategy.png','4.5  |  Evidence-to-action implementation roadmap',[
 ('INSTRUMENT','Existing event helper\nAllowlisted metadata\nNo keys in GA4'),
 ('REDUCE FRICTION','Existing demo connect\nProvider fallback\nLocal dashboard flow'),
 ('VALIDATE VALUE','Proposed real-session\nactivation event\nMature cohort checks'),
 ('EXPERIMENT','Proposed checklist\nWeekly insights digest\nRetention guardrails')],
 'Status: first two steps exist in code; later steps are recommendations, not completed deployments', [TEAL,TEAL,ORANGE,ORANGE])
# Figure 4.6
fig,ax=plt.subplots(figsize=(12,5))
names=['Visitors','Signups','Key configurations','Monitoring sessions']; ns=[len(visitors),len(users),len(configs),len(sessions)]
ax.barh(names[::-1],ns[::-1],height=.55,color=[NAVY,BLUE,ORANGE,TEAL][::-1])
for i,v in enumerate(ns[::-1]):ax.text(v+240,i,f'{v:,}',va='center')
ax.set_xlim(0,25000);ax.set_xlabel('Rows in synthetic primary datasets')
finish(fig,'4_6_synthetic_data.png','4.6  |  Reproducible synthetic dataset inventory','Seed 42 | Generated scenario, not collected beta-user activity | Original datasets unchanged')
# Figure 4.7
fig,ax=plt.subplots(figsize=(12,5))
ax.set_xlim(0,12);ax.set_ylim(0,5);ax.axis('off')
boxes=[(.2,1.55,2.8,1.7,'RANDOMIZE','New signups | 50:50\nStratify by location',NAVY),
       (4.25,2.85,3.5,1.5,'CONTROL A','Current onboarding\nSame demo access',BLUE),
       (4.25,.5,3.5,1.65,'TREATMENT B','Provider guidance + checklist\nOne follow-up nudge',TEAL),
       (9,1.55,2.8,1.7,'EVALUATE','Real session within 7 days\n95% CI + guardrails',ORANGE)]
for x,y,w,h,heading,body,c in boxes:
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.07,rounding_size=0.12',facecolor='#edf5fa',edgecolor=c,linewidth=1.6))
    ax.text(x+w/2,y+h-.37,heading,ha='center',fontweight='bold',color=c,fontsize=12)
    ax.text(x+w/2,y+.56,body,ha='center',va='center',fontsize=10,linespacing=1.5)
for start,end in [((3.1,2.7),(4.1,3.5)),((3.1,2.1),(4.1,1.3)),((7.9,3.5),(8.85,2.7)),((7.9,1.3),(8.85,2.1))]:
    ax.annotate('',xy=end,xytext=start,arrowprops={'arrowstyle':'->','color':GRAY,'lw':2})
finish(fig,'4_7_ab_test.png','4.7  |  Proposed onboarding A/B test','EXPERIMENT DESIGN ONLY | No measured outcomes, significance or uplift | Participants enter one arm only')
# Figure 4.8
diagram('4_8_model_training.png','4.8  |  Existing ML design and validation gate',[
 ('INPUTS','Prompt + response\nTokens + latency\nProvider metadata'),
 ('FEATURES','Text / pattern counts\nLeakage indicators\nEngineered ratios'),
 ('XGBOOST','Six risk regressors\nOptional CRI classifier\n100 trees; depth 10'),
 ('VALIDATE FIRST','Fix label consistency\nGroup-held-out split\nMAE, recall, latency')],
 'Code architecture, not a completed benchmark | No trained artifact or held-out score is claimed', [NAVY,BLUE,TEAL,ORANGE])
# Figure 4.9
fig,ax=plt.subplots(figsize=(12,5))
risks=['Injection','Leakage','Hallucination','Bias','Anomaly','Tool misuse'];weights=[.25,.20,.15,.10,.15,.15]
ax.barh(risks[::-1],weights[::-1],color=TEAL,height=.6)
for i,v in enumerate(weights[::-1]):ax.text(v+.006,i,f'{v:.0%}',va='center')
ax.set_xlim(0,.33);ax.set_xlabel('Configured contribution to weighted Cognitive Risk Index')
finish(fig,'4_9_risk_weights.png','4.9  |  From predicted risks to CRI recommendations','Source: etl/ml-service/config.py | Configuration weights, not measured model predictions')
# Figure 4.10
fig,ax=plt.subplots(figsize=(12,5))
complete=[r for r in monthly if r['month']<'2026-09']
ax.plot([r['month_label'] for r in complete],[int(r['mau']) for r in complete],marker='o',lw=2.5,color=TEAL)
for i,r in enumerate(complete):ax.annotate(r['mau'],(i,int(r['mau'])),xytext=(0,10),textcoords='offset points',ha='center')
ax.set_ylim(0,500);ax.set_ylabel('Monthly active users (observed sessions)');ax.grid(axis='y',alpha=.15)
finish(fig,'4_10_results.png','4.10  |  Simulated engagement across complete months','Synthetic scenario | September excluded because it ends on 20 Sep | Trend is not causal evidence')

# Report content: one shared source for Markdown, HTML, Word, and PDF.
sections=[]
def section(title,blocks): sections.append((title,blocks))
def p(t):return ('p',t)
def bullets(*items):return ('bullets',items)
def table(headers,rows):return ('table',(headers,rows))
def image(name,caption):return ('image',(name,caption))

section('4.1 Funnel Analysis',[
p('Funnel analysis examines how users move from discovering monitr-ai to receiving monitoring value. The supplied synthetic data covers 1 March to 20 September 2026. The funnel uses unique visitors and unique users; multiple API-key configuration rows must not be counted as separate users.'),
table(['Stage','Count','Conversion from previous stage'],[
 ['Website visit','20,731','100.00% (starting population)'],['Account signup','1,440','6.95%'],['API key connected','964','66.94%'],['Activation flag','813','84.34%']]),
p('The largest absolute drop is visitor to signup: 19,291 visitors do not create an account. After signup, 476 users do not connect a key, and 151 connected users lack an activation flag. End-to-end flagged activation is 3.92% of visitors; 56.46% of signups receive the activation flag. The dataset reports a one-day median signup-to-configuration delay.'),
p('Interpretation: improve landing-page value communication and the provider-connection experience. These are hypotheses supported by simulated drop-offs, not proven reasons for abandonment. The generator defines activation within seven days of key connection, not signup. Importantly, only 804 of the 813 flagged users have any session record: nine synthetic flags are inconsistent with observed activation. Therefore the report calls this an activation-flag funnel, not a verified first-session funnel.'),
image('4_1_funnel.png','Figure 4.1. Synthetic acquisition funnel. Counts use the supplied flags, including the nine flagged users without a recorded session.'),
p('Evidence: analytics/data/funnel_overall.csv; user_signups.csv; monitoring_sessions.csv; analytics/generate_data.py (signup and session generation).')])

section('4.2 Churn Analysis',[
p('monitr-ai is described as a pre-launch, free-beta product. Subscription cancellations, paid churn and revenue churn are therefore not appropriate measures. This analysis uses the existing inactivity proxy: an activated user with more than 30 days since the last recorded monitoring session as of 20 September 2026.'),
table(['Segment','Inactive / activation flags','Inactivity rate'],[['Overall','295 / 813','36.3%'],['DSU Campus','81 / 226','35.8%'],['Outside Campus','214 / 587','36.5%']]),
p('The campus and outside-campus inactivity rates differ by only 0.7 percentage points. The data does not support a claim that campus users churn substantially more. Outside-campus users activate more often, but their snapshot inactivity is not lower. Usage-history length, recent activation and incident-driven use may affect this proxy; the snapshot denominator is not exposure-adjusted and the nine flagged users without sessions are not identified as inactive by the generator.'),
p('A separate retention measure checks whether a user has a session in activation days [28,35). To avoid right-censoring, this report requires the full 35-day observation window. In the August signup cohort, 46 of 63 eligible users return in week four (73.0%). The supplied monthly CSV instead reports 48.4%, using an incompletely observed cohort denominator. These values answer different calculations and must not be silently interchanged. August-to-September return of 77.4% is also provisional because September stops on the 20th.'),
image('4_2_inactivity.png','Figure 4.2. Snapshot inactivity rates using the existing all-activation-flags denominator.'),
p('Recommended retention actions: opt-in weekly insight summaries, meaningful CRI alerts and first-session support. Evaluate them using mature cohorts rather than treating a synthetic correlation as causation.')])

section('4.3 Data Generators & Import',[
p('The repository has two distinct data-generation paths: a product-analytics generator and a security-training generator. analytics/generate_data.py creates linked CSVs with random seed 42. These files feed the Tableau analytics pack and the figures in this report. They are not exports of real visitors or real campus participants.'),
table(['Dataset / component','Purpose','Rows / mechanism'],[['web_visitors.csv','Acquisition and visitor-to-signup joins','20,731 rows'],['user_signups.csv','User attributes, onboarding and flags','1,440 rows'],['api_key_configs.csv','Provider configuration attempts / records','1,194 rows; 964 users'],['monitoring_sessions.csv','Usage, quality and safety telemetry','13,385 rows'],['users_dsu_campus / users_outside_campus','Location partitions of signups','452 + 988 = 1,440']]),
p('For the intended live pipeline, Convex exports telemetry to etl/webhook.ts, which writes to PostgreSQL. dbt staging models standardize raw events; marts aggregate daily metrics, provider comparisons and cognitive waste. The frontend GA4 helper records allowlisted onboarding metadata. The current key-connection demo saves configuration locally, so it should not be described as a verified live API-key validation import. CSV analytics are a separate offline analysis path.'),
p('Import validation in this report checks unique visitor/user/session IDs, session-to-user and session-to-config references, and agreement between user session totals and 13,385 session rows. These checks pass. A separate semantic check identifies the nine activation flags without sessions. Before live ingestion, also validate timestamps, duplicate event IDs, redaction, schema versions and webhook authentication.'),
image('4_3_import_pipeline.png','Figure 4.3. Conceptual data architecture. It does not assert that external services were deployed or contacted.'),
p('Evidence: analytics/generate_data.py; analytics/build_workbook.py; convex/etlExport.ts; etl/webhook.ts; etl/models/; frontend/analytics.js.')])

section('4.4 Analysis and Identifying the Issue(s) and Solution(s) for the Same (With Data Evidence)',[
p('The following issues are grounded in the supplied synthetic records or inspected code. Proposed causes remain hypotheses until tested on observed product activity.'),
table(['Issue and evidence','Interpretation / proposed solution'],[
 ['Only 6.95% of visitors sign up; 19,291 drop off.','Test clearer value propositions, sample insights and a concise demo CTA.'],
 ['476 of 1,440 signups lack a connected key (33.1%).','Offer provider guidance and an explicit demo path; measure actual connection success separately.'],
 ['Campus activation 50.0% vs outside-campus 59.4%: a 9.4-point gap.','Test a workshop follow-through checklist; do not infer that location causes low activation.'],
 ['Campus inactivity 35.8% vs outside-campus 36.5%.','Treat inactivity as a shared issue; do not claim disproportionate campus churn.'],
 ['August CWI 0.429 vs the repo target <=0.25; efficiency 0.774 vs >=0.80.','Investigate prompt length, repeated context and provider/model fit; validate changes against quality.'],
 ['Nine activation flags have no session; immature retention cohorts enter some denominators.','Derive observed activation from sessions and require complete retention windows.'],
 ['Synthetic security training repeats benign prompts and uses template-derived labels.','Deduplicate; hold out attack families; obtain independently reviewed labels before deployment.']]),
image('4_4_segment_evidence.png','Figure 4.4. Location activation and inactivity use different denominators. Location differences are descriptive, not causal.'),
p('Targets are internal values from kpi_definitions.csv, not externally validated industry benchmarks. Its prefilled status labels are not reliable evidence of target attainment; the numeric comparison above takes precedence.')])

section('4.5 Arrive at Strategy and Implement',[
p('The recommended strategy is to shorten time-to-value, make repeated monitoring useful, and validate security scoring before enforcement. The repository already contains parts of this strategy, but code presence is not proof of production deployment or improved outcomes.'),
table(['Workstream','Current evidence','Next implementation / acceptance criterion'],[
 ['Instrument onboarding','analytics.js records key-entry, submission, acceptance and dashboard events with safe metadata.','Add a server-confirmed first monitoring-session event, deduplication and a defined activation window.'],
 ['Remove demo friction','key.html has DEMO_MODE, provider fallback and local dashboard navigation.','Keep demo vs validated-provider success separate; verify a real provider only in the production flow.'],
 ['Enable recurring value','Feature flags show 65.9% synthetic custom-alert adoption.','Test an opt-in weekly insights digest and useful threshold alerts; monitor opt-outs and false alerts.'],
 ['Make risk scoring safe','ML service contains six-risk scoring and CRI-based actions.','Use shadow mode and independent labels before allowing automated restrictions.']]),
p('Priority order: repair metric definitions first; validate the first-session event next; then test onboarding and retention changes. Product analytics owners should track activation, while security owners audit model errors and threshold decisions. This report and ten reproducible figures are implemented deliverables; new product features and live experiments are not claimed as completed.'),
image('4_5_strategy.png','Figure 4.5. Strategy roadmap explicitly distinguishes existing code from recommended work.'),
p('Privacy requirement: never place API-key material in analytics events or reports. Existing GA4 safeguards allowlist parameter names and filter secret-looking values. Real credentials should not rely on the demo localStorage mechanism for production security.')])

section('4.6 Synthetic Data Generation',[
p('Synthetic data enables reproducible analytics demonstrations without exposing customer activity. The product generator fixes random.seed(42), uses dates from 1 March to 20 September 2026, and simulates acquisition channels, personas, campus/outside-campus segments, provider setups, usage lifetimes, telemetry and feature adoption.'),
p('Primary join paths are anonymous_visitor_id to visitor_id, user_id across signup/config/session tables, and config_id from sessions to configurations. Location-specific user CSVs partition the signup table and must be unioned rather than counted as extra users. Derived funnel, monthly KPI, cohort and feature-adoption CSVs summarize these records.'),
p('A second generator in etl/ml-service/processor.py builds security-training examples: five benign prompt templates repeated ten times (50 examples), plus the first three phrases from each of nine attack categories (27 examples). This produces 77 rows, many of which are duplicate or closely related examples. Risk labels are hand-assigned templates, not expert judgments of real attacks.'),
p('Generation limitations matter: location, persona and onboarding effects are deliberately encoded in probabilities. Their apparent relationships are simulation assumptions, not discovered behavioral laws. Security labels mainly cover low-risk ranges; they do not adequately validate high or critical risk. Expand scenarios, diversify providers and language, and split by template family before using this data for model evaluation.'),
image('4_6_synthetic_data.png','Figure 4.6. Row counts of the primary product-analytics datasets. These are not human participant counts.'),
p('Rebuild original data only when needed: python3 analytics/generate_data.py; python3 analytics/build_workbook.py. This report reads existing CSVs and does not regenerate or overwrite the original analytics pack.')])

section('4.7 A/B Test',[
p('No recorded randomized experiment, treatment assignments or A/B outcomes were found in the inspected project. This section presents a proposed experiment, not a completed test. Hypothesis: guided onboarding with a first-session checklist increases observed activation compared with the current flow.'),
table(['Design element','Proposed specification'],[
 ['Unit and assignment','New signup; stable 50:50 assignment, stratified by collection location; exclude internal/test accounts.'],
 ['Control A','Current onboarding flow.'],
 ['Treatment B','Same flow plus provider guidance, a first-session checklist and one follow-up nudge.'],
 ['Primary outcome','At least one real monitoring session within seven days of signup; intention-to-treat denominator.'],
 ['Secondary / guardrails','Connection success, session quality, mature W4 retention, setup errors, opt-outs and safety alerts.'],
 ['Decision rule','Pre-register sample size and analysis; report absolute uplift and 95% confidence interval after all follow-up windows close.']]),
p('Illustrative planning calculation only: using the synthetic flagged signup activation rate p1 = 813/1,440 = 56.46% as a rough baseline and a minimum detectable absolute improvement of 8 percentage points (p2 = 64.46%), a two-sided normal approximation with alpha 0.05 and 80% power requires about 586 signups per arm (1,172 total). Re-estimate this after measuring the real primary outcome, since the synthetic baseline uses a different activation definition. Do not stop early for significance without a sequential-testing design.'),
p('Sample-size formula per arm: n = [1.96*sqrt(2*p_bar*(1-p_bar)) + 0.842*sqrt(p1*(1-p1)+p2*(1-p2))]^2 / (p2-p1)^2. No uplift, p-value, confidence interval or winner is reported because no experimental observations exist.'),
image('4_7_ab_test.png','Figure 4.7. Proposed experiment flow. This is not a chart of observed A/B performance.')])

section('4.8 Identify the Right Model and Train',[
p('The repository\'s ML task is event-level security-risk estimation, not user-churn prediction. etl/ml-service/model.py implements MultiOutputRegressor around XGBRegressor to estimate injection, leakage, hallucination, bias, anomaly and tool-misuse risk. An additional XGBClassifier predicts CRI levels. XGBoost is a plausible candidate for engineered tabular features, but it has not been demonstrated to be the best model by this report.'),
p('Feature extraction includes prompt/response lengths, attack-pattern indicators, leakage markers, token/latency ratios and provider/model metadata. Default model settings are 100 estimators, maximum depth 10, learning rate 0.1 and random state 42. The fit routine scales features and trains the regressors/classifier. A scaler is not generally required for tree models, but it is present in the code.'),
table(['Validation gate','Required action before claiming a trained, reliable model'],[
 ['Consistent labels','Replace the synthetic trainer\'s unweighted mean CRI labels with labels based on the configured weighted CRI, or document an independently reviewed target.'],
 ['Coverage and class handling','The 77-row generator supplies minimal/low labels, while the classifier config specifies five classes. Test library compatibility and class handling; add high-risk coverage.'],
 ['Leakage-resistant split','Deduplicate repeated benign prompts; hold out prompt/attack families or later time periods before fitting preprocessing.'],
 ['Candidate comparison','Compare constant/rule baselines, regularized linear regression, random forest and XGBoost using the same folds.'],
 ['Evaluation and audit','Report per-risk MAE/RMSE, high-risk recall, false-positive rates, calibration and inference latency on independent labels.']]),
p('The inspected BatchTrainer sends its available examples directly into fit; it does not execute a held-out evaluation. Training entry points are python train.py --synthetic and python train.py --database from etl/ml-service. They require service dependencies and, for the database path, configured labeled data. This report does not run training, certify these commands, create a model artifact or invent accuracy scores.'),
image('4_8_model_training.png','Figure 4.8. Existing model design with proposed validation gates, not a completed model-selection result.')])

section('4.9 Model Predictions and Recommendations',[
p('At inference, the intended service extracts features, predicts six risk values, clips each to [0,1], and computes a weighted Cognitive Risk Index (CRI). The configured equation is CRI = 0.25*injection + 0.20*leakage + 0.15*hallucination + 0.10*bias + 0.15*anomaly + 0.15*tool_misuse.'),
table(['CRI range','Configured action','Recommended operational handling'],[
 ['0.00 <= CRI < 0.20','allow','Log and monitor; a low score is not a safety guarantee.'],
 ['0.20 <= CRI < 0.50','warn','Show risk context; review flagged content if relevant.'],
 ['0.50 <= CRI < 0.75','restrict','Consider limiting tools or sensitive operations after validation.'],
 ['0.75 <= CRI <= 1.00','block','Candidate for blocking, with an audited override / review path.']]),
p('Worked arithmetic example, not model output: risk vector [0.80, 0.30, 0.20, 0.10, 0.40, 0.20] gives CRI = 0.390. The configured action is warn; the separate trust mapping marks it restricted (trust threshold 0.25). CRI-level category thresholds differ from action thresholds, and the optional classifier can disagree with the weighted-score category. Report these fields separately rather than assuming they are interchangeable.'),
p('Recommended user-facing responses include isolating untrusted instructions for injection risk, redacting sensitive output for leakage risk, and checking claims against trusted evidence for hallucination risk. A lexical feature model cannot verify factual truth or prove an interaction safe. Begin in shadow mode; calibrate thresholds and assess false positives before any automated enforcement.'),
image('4_9_risk_weights.png','Figure 4.9. Configured CRI risk weights. No live or trained-model predictions are presented.'),
p('Evidence: etl/ml-service/config.py; model.py (_compute_cri, _get_action, _get_trust_level, predict); main.py prediction endpoints.')])

section('4.10 Observations/Results',[
p('The reproducible result is a descriptive analysis of the existing synthetic project dataset, supported by source inspection. It establishes where the simulated funnel leaks and which measurements need correction; it does not establish real-world business impact.'),
table(['Observation','Result and interpretation'],[
 ['Dataset scale','20,731 visitors; 1,440 signups; 1,194 key configs; 13,385 sessions.'],
 ['Activation','813 flagged users (56.46% of signups), but 804 session-observed users; nine flags need reconciliation.'],
 ['Inactivity','295 / 813 = 36.3% using the repository proxy; not cancellation churn.'],
 ['Location','50.0% campus vs 59.4% outside-campus activation; no meaningful descriptive campus inactivity excess.'],
 ['Complete-month engagement','Synthetic MAU grows from 68 in March to 429 in August; August DAU/MAU is 20.1%.'],
 ['Quality / safety (August)','CWI 0.429; efficiency 0.774; average CRI 0.205; these are simulated telemetry aggregates.'],
 ['Experiment and ML outcomes','No observed A/B uplift, validated model accuracy or independently measured deployment impact is available.']]),
image('4_10_results.png','Figure 4.10. Synthetic MAU trend across complete months. Increased simulated usage is not evidence that an intervention worked.'),
p('Conclusion: prioritize clean observed-activation metrics, guided onboarding and validated risk scoring. Use mature cohorts for retention and a randomized test to estimate intervention effects. Keep product-analytics simulation, model training, and real production measurements clearly separated.'),
p('Evidence limitations: no live GA4, Firebase, Convex or PostgreSQL data was accessed; no experiment or model training was executed. The original analytics CSVs and application implementation remain unchanged. All ten figures were programmatically created from repository values or labeled implementation/design diagrams.')])

intro = [
 'Project: monitr-ai - LLM Observability Platform',
 'Chapter 4 | Implementation',
 'Prepared 23 October 2026 | Dataset window: 1 March - 20 September 2026',
 'Evidence notice: All supplied product-analytics records are SYNTHETIC. Charts are not real customer results. Existing code, descriptive calculations, illustrative arithmetic and proposed work are identified separately. No A/B uplift or trained-model performance is invented.',
 'This chapter addresses Sections 4.1-4.10 with ten reusable figures. It is based on the repository README, analytics CSVs and generator, frontend event/demo code, ETL implementation and ML-service source.'
]
md=['# '+intro[1], '\n## '+intro[0], '\n'+intro[2], '\n> '+intro[3], '\n'+intro[4], '\n## Contents']
md += ['- '+title for title,_ in sections]
for title,blocks in sections:
    md+=['\n## '+title+'\n']
    for kind,data in blocks:
        if kind=='p':md += [data+'\n']
        elif kind=='bullets':md += ['\n'.join('- '+t for t in data)+'\n']
        elif kind=='table':
            headers,rows=data
            md += ['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']
            md += ['| '+' | '.join(row)+' |' for row in rows];md += ['']
        elif kind=='image':
            name,caption=data
            md += [f'![{caption}](images/{name})\n',f'*{caption}*\n']
md += ['\n## Source files and reproduction\n',
       'Primary sources: `README.md`, `analytics/README.md`, `analytics/summary.json`, `analytics/data/*.csv`, `analytics/generate_data.py`, `frontend/analytics.js`, `frontend/key.html`, `convex/etlExport.ts`, `etl/webhook.ts`, `etl/models/`, and `etl/ml-service/{config,features,model,processor,train,main}.py`.\n',
       'Reproduce this chapter: `python3 -m venv .venv`, `.venv/bin/pip install -r "product analytics/requirements.txt"`, then `.venv/bin/python "product analytics/build_report.py"`. Original source datasets are read-only inputs. `evidence.json` records computed counts and mature week-four retention.\n']
mdtext='\n'.join(md)
(OUT/'implementation.md').write_text(mdtext)
html='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>monitr-ai | Implementation</title><style>
body{font-family:system-ui,sans-serif;max-width:1000px;margin:48px auto;padding:0 24px;color:#142b49;line-height:1.65;background:#fbfdff}h1{font-size:40px}h2{border-bottom:2px solid #008c95;padding-bottom:12px;margin-top:48px}img{width:100%;height:auto;border:1px solid #e0e8f0;border-radius:8px}table{border-collapse:collapse;width:100%;font-size:14px;margin:24px 0}th{background:#142b49;color:white}th,td{padding:12px;text-align:left;border:1px solid #dfe6ed}tr:nth-child(even){background:#eff5fa}blockquote{border-left:5px solid #df8735;background:#fff4e8;padding:16px 24px;margin:24px 0}code{overflow-wrap:anywhere}@media print{h2{break-before:page}img{break-inside:avoid}body{margin:0;max-width:none}}
</style><main>'''+markdown.markdown(mdtext,extensions=['tables'])+'</main></html>'
(OUT/'implementation.html').write_text(html)

# Word document: page per requested subsection, fully editable text and tables.
doc=Document()
doc.sections[0].header.paragraphs[0].text='monitr-ai  /  IMPLEMENTATION  /  SYNTHETIC-DATA STUDY'
normal=doc.styles['Normal']; normal.font.name='Calibri';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(7)
doc.add_heading(intro[1],0);doc.add_heading(intro[0],1)
for t in intro[2:]:doc.add_paragraph(t)
doc.add_heading('Contents',2)
for title,_ in sections:doc.add_paragraph(title)
for title,blocks in sections:
    doc.add_page_break();doc.add_heading(title,1)
    for kind,data in blocks:
        if kind=='p':doc.add_paragraph(data)
        elif kind=='bullets':
            for item in data:doc.add_paragraph(item,style='List Bullet')
        elif kind=='table':
            headers,rows=data;t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
            for i,v in enumerate(headers):t.rows[0].cells[i].text=v
            for row in rows:
                cells=t.add_row().cells
                for i,v in enumerate(row):cells[i].text=v
        elif kind=='image':
            name,caption=data;doc.add_picture(str(IMG/name),width=Inches(6.1));doc.add_paragraph(caption,style='Caption')
doc.add_heading('Reproduction',2);doc.add_paragraph('See build_report.py, requirements.txt, evidence.json and implementation.md in the same directory. All generated images are in images/.')
doc.save(OUT/'implementation.docx')

# PDF, same content, with tables wrapped rather than clipped.
styles=getSampleStyleSheet()
styles['Normal'].fontName='Helvetica';styles['Normal'].fontSize=9;styles['Normal'].leading=12
styles['Heading1'].textColor=colors.HexColor(NAVY);styles['Heading1'].fontSize=17;styles['Heading1'].leading=21
styles['Heading2'].textColor=colors.HexColor(TEAL)
styles.add(styles['Normal'].clone('Caption'));styles['Caption'].fontSize=8;styles['Caption'].leading=10;styles['Caption'].textColor=colors.HexColor(GRAY)
styles.add(styles['Normal'].clone('Cell'));styles['Cell'].fontSize=8;styles['Cell'].leading=10
story=[Paragraph(escape(intro[1]),styles['Title']),Spacer(1,20),Paragraph(escape(intro[0]),styles['Heading1'])]
for t in intro[2:]:story += [Spacer(1,12),Paragraph(escape(t),styles['Normal'])]
story += [Spacer(1,24),Paragraph('Contents',styles['Heading2'])]
for title,_ in sections:story += [Paragraph(escape(title),styles['Normal']),Spacer(1,5)]
for title,blocks in sections:
    story += [PageBreak(),Paragraph(escape(title),styles['Heading1']),Spacer(1,8)]
    for kind,data in blocks:
        if kind=='p':story += [Paragraph(escape(data),styles['Normal']),Spacer(1,9)]
        elif kind=='bullets':
            for item in data:story += [Paragraph('- '+escape(item),styles['Normal'])]
        elif kind=='table':
            headers,rows=data
            vals=[[Paragraph(escape(v),styles['Cell']) for v in row] for row in [headers]+rows]
            widths=[170,342] if len(headers)==2 else [150,155,207]
            t=Table(vals,colWidths=widths,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dceaf3')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f2f6fa')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor(TEAL))]))
            story += [t,Spacer(1,10)]
        elif kind=='image':
            name,caption=data
            story += [Image(str(IMG/name),width=512,height=512*5/12),Spacer(1,4),Paragraph(escape(caption),styles['Caption']),Spacer(1,10)]
def footer(canvas,doc):
    canvas.setStrokeColor(colors.HexColor('#dce4ee'));canvas.line(42,35,554,35)
    canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor(GRAY))
    canvas.drawString(42,23,'monitr-ai | Synthetic-data analysis | Prepared 23 Oct 2026')
    canvas.drawRightString(554,23,str(doc.page))
SimpleDocTemplate(str(OUT/'implementation.pdf'),pagesize=(595.28,841.89),leftMargin=42,rightMargin=41,topMargin=40,bottomMargin=45,title='monitr-ai - Chapter 4 Implementation',author='monitr-ai project documentation').build(story,onFirstPage=footer,onLaterPages=footer)

# Machine-readable audit and portable bundle.
p1=len(activated)/len(users);p2=p1+.08;pbar=(p1+p2)/2
sample=math.ceil((1.95996398454*math.sqrt(2*pbar*(1-pbar))+.84162123357*math.sqrt(p1*(1-p1)+p2*(1-p2)))**2/(p2-p1)**2)
evidence={'data_kind':'synthetic','as_of':ASOF.isoformat(),'visitors':len(visitors),'signups':len(users),'config_rows':len(configs),'connected_users':len({c['user_id'] for c in configs}),'activation_flags':len(activated),'session_observed_users':observed,'activation_flags_without_sessions':missing,'sessions':len(sessions),'inactive_flag_count':len(inactive),'inactivity_denominator':len(activated),'inactivity_pct':round(len(inactive)/len(activated)*100,2),'mature_w4_by_signup_cohort':w4,'ab_test_status':'proposed, not executed','ab_planning_signups_per_arm':sample,'model_training_status':'source inspected; not executed or independently validated','checks':{'unique_primary_ids':True,'session_user_and_config_references':True,'session_totals_reconcile':True},'figures':10}
(OUT/'evidence.json').write_text(json.dumps(evidence,indent=2)+'\n')
with zipfile.ZipFile(OUT/'implementation_pack.zip','w',compression=zipfile.ZIP_DEFLATED) as z:
    for f in sorted(OUT.rglob('*')):
        if f.is_file() and f.suffix!='.zip' and '__pycache__' not in f.parts:
            z.write(f,f.relative_to(OUT))
print(json.dumps(evidence,indent=2))
print('Built Markdown, HTML, DOCX, PDF, ten PNG figures and ZIP in',OUT)
